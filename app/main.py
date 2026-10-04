from __future__ import annotations

import asyncio
import csv
import io
import math
import os
import re
import sqlite3
import zipfile
from contextlib import asynccontextmanager
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Annotated, Any

import httpx
from fastapi import Depends, FastAPI, File, Header, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel
from starlette.concurrency import run_in_threadpool
from starlette.middleware.base import BaseHTTPMiddleware

BASE_DIR = Path(__file__).resolve().parent.parent
STATIC_DIR = BASE_DIR / "static"
DATA_DIR = Path(os.getenv("DATA_DIR", BASE_DIR / "data"))
DB_PATH = DATA_DIR / "merchants.db"
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024

IMPORT_COLUMNS = ("사업자번호", "가맹점번호", "가맹점명", "사업장주소", "가맹점업종명")
DISCARDED_PII_COLUMNS = {"대표자명", "휴대전화번호"}
KNOWN_SOURCE_COLUMNS = {
    "지역구분", "그룹가맹점번호", "사업자번호", "가맹점번호", "가맹점명",
    "지역가맹점 등록여부", "대표자명", "휴대전화번호", "사업장주소", "등록일자",
    "등록부점코드", "등록텔러번호", "해지일자", "해지부점코드", "해지직원번호",
    "가맹점우편번호", "우편주소", "가맹점상세주소", "가맹점업종명",
    "가맹점신규일자", "체크카드가맹점수수료율", "사업자 상태", "폐업일",
}
PII_PATTERNS = {
    "주민등록번호": re.compile(r"(?<!\d)(?:\d{6})[-\s]?[1-4]\d{6}(?!\d)"),
    "휴대전화번호": re.compile(r"(?<!\d)01[016789][-.\s]?\d{3,4}[-.\s]?\d{4}(?!\d)"),
    "유선전화번호": re.compile(r"(?<!\d)0(?:2|[3-6][1-5])[-.\s]?\d{3,4}[-.\s]?\d{4}(?!\d)"),
    "이메일주소": re.compile(r"(?<![\w.])[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}(?![\w.])"),
}

SAMPLE_MERCHANTS = (
    ("1010100001", "DGP-001", "동명정육점", "광주 동구 동명로 18", "축산물", 35.14935, 126.92355),
    ("1010100002", "DGP-002", "푸른길 카페", "광주 동구 동계천로 137", "카페·디저트", 35.14996, 126.92512),
    ("1010100003", "DGP-003", "산수우리약국", "광주 동구 필문대로 184", "의료·약국", 35.15472, 126.93042),
    ("1010100004", "DGP-004", "대인시장 반찬가게", "광주 동구 제봉로194번길 7", "식품", 35.15468, 126.91338),
    ("1010100005", "DGP-005", "충장서림", "광주 동구 충장로 87", "도서·문구", 35.14821, 126.91485),
    ("1010100006", "DGP-006", "무등세탁소", "광주 동구 경양로 342", "생활서비스", 35.15916, 126.92317),
    ("1010100007", "DGP-007", "계림떡방앗간", "광주 동구 무등로 321", "식품", 35.16028, 126.91832),
    ("1010100008", "DGP-008", "학동우리분식", "광주 동구 남문로 689", "음식점", 35.13388, 126.92735),
    ("1010100009", "DGP-009", "조대안경원", "광주 동구 필문대로 309", "안경", 35.14285, 126.92992),
    ("1010100010", "DGP-010", "금남로 꽃집", "광주 동구 금남로 203", "화훼", 35.14776, 126.91873),
    ("1010100011", "DGP-011", "동구청앞 편의점", "광주 동구 서남로 1", "편의점", 35.14591, 126.92316),
    ("1010100012", "DGP-012", "문화전당 공방", "광주 동구 문화전당로 38", "공예·잡화", 35.14639, 126.91991),
)


class Merchant(BaseModel):
    id: int
    name: str
    address: str
    category: str
    lat: float
    lng: float
    distance: float | None = None
    business_status: str | None = None


class MerchantList(BaseModel):
    items: list[Merchant]
    total: int


class SecurityHeadersMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        response = await call_next(request)
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "SAMEORIGIN"
        response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
        response.headers["Permissions-Policy"] = "geolocation=(self)"
        return response


def db_connect() -> sqlite3.Connection:
    connection = sqlite3.connect(DB_PATH, timeout=30)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA journal_mode = WAL")
    return connection


def initialize_database() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    with db_connect() as connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS merchants (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                business_no TEXT NOT NULL,
                merchant_no TEXT NOT NULL UNIQUE,
                name TEXT NOT NULL,
                address TEXT NOT NULL,
                category TEXT NOT NULL,
                lat REAL,
                lng REAL,
                business_status TEXT,
                status_checked_at TEXT,
                uploaded_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_merchants_location ON merchants(lat, lng);
            CREATE INDEX IF NOT EXISTS idx_merchants_category ON merchants(category);
            CREATE TABLE IF NOT EXISTS upload_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                filename TEXT NOT NULL,
                row_count INTEGER NOT NULL,
                geocode_failed_count INTEGER NOT NULL,
                uploaded_at TEXT NOT NULL
            );
            """
        )
        count = connection.execute("SELECT COUNT(*) FROM merchants").fetchone()[0]
        if count == 0 and os.getenv("LOAD_SAMPLE_DATA", "true").lower() == "true":
            now = datetime.now(UTC).isoformat()
            connection.executemany(
                """INSERT INTO merchants
                   (business_no, merchant_no, name, address, category, lat, lng, business_status, uploaded_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, '계속사업자', ?)""",
                [(*merchant, now) for merchant in SAMPLE_MERCHANTS],
            )


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", normalize_text(value).lstrip("\ufeff"))


def validate_headers(headers: list[str]) -> tuple[list[int], dict[str, int]]:
    if not any(headers):
        raise ValueError("첫 번째 행에 컬럼명이 필요합니다.")
    named_headers = [header for header in headers if header]
    if len(named_headers) != len(set(named_headers)):
        raise ValueError("중복된 컬럼명이 있습니다.")
    missing = [column for column in IMPORT_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"필수 컬럼이 없습니다: {', '.join(missing)}")
    unknown = [header for header in headers if header and header not in KNOWN_SOURCE_COLUMNS]
    if unknown:
        raise ValueError(f"허용되지 않은 컬럼이 있습니다: {', '.join(unknown[:5])}")
    safe_scan_indexes = [
        index for index, header in enumerate(headers)
        if header and header not in DISCARDED_PII_COLUMNS
    ]
    indexes = {column: headers.index(column) for column in IMPORT_COLUMNS}
    if "사업자 상태" in headers:
        indexes["사업자 상태"] = headers.index("사업자 상태")
    return safe_scan_indexes, indexes


def parse_values(
    values: list[str],
    row_number: int,
    headers: list[str],
    safe_scan_indexes: list[int],
    import_indexes: dict[str, int],
    merchant_numbers: set[str],
) -> dict[str, str] | None:
    item = {column: values[index] if index < len(values) else "" for column, index in import_indexes.items()}
    if not any(item.get(column) for column in IMPORT_COLUMNS):
        return None
    for index in safe_scan_indexes:
        value = values[index] if index < len(values) else ""
        for pii_name, pattern in PII_PATTERNS.items():
            if pattern.search(value):
                raise ValueError(f"{row_number}행 {headers[index]}에서 {pii_name} 패턴이 탐지되었습니다.")
    if "폐업" in item.get("사업자 상태", ""):
        return None
    empty_required = [column for column in IMPORT_COLUMNS if not item[column]]
    if empty_required:
        raise ValueError(f"{row_number}행 필수값이 비어 있습니다: {', '.join(empty_required)}")
    item["사업자번호"] = re.sub(r"\D", "", item["사업자번호"])
    if len(item["사업자번호"]) != 10:
        raise ValueError(f"{row_number}행 사업자번호는 숫자 10자리여야 합니다.")
    merchant_no = item["가맹점번호"]
    if merchant_no in merchant_numbers:
        raise ValueError(f"{row_number}행 가맹점번호가 중복되었습니다: {merchant_no}")
    merchant_numbers.add(merchant_no)
    return item


def validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1_000 or sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("압축 해제 크기가 허용 범위를 초과합니다.")
            if any("vbaProject.bin" in entry.filename for entry in entries):
                raise ValueError("매크로가 포함된 파일은 업로드할 수 없습니다.")
    except zipfile.BadZipFile as exc:
        raise ValueError("올바른 .xlsx 파일이 아닙니다.") from exc


def parse_excel(content: bytes) -> list[dict[str, str]]:
    validate_xlsx_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("올바른 .xlsx 파일이 아닙니다.") from exc

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=False)
        try:
            header_cells = next(rows)
        except StopIteration as exc:
            raise ValueError("엑셀 파일이 비어 있습니다.") from exc

        headers = [normalize_header(cell.value) for cell in header_cells]
        safe_scan_indexes, import_indexes = validate_headers(headers)
        parsed: list[dict[str, str]] = []
        merchant_numbers: set[str] = set()

        for row_number, cells in enumerate(rows, start=2):
            values = [normalize_text(cell.value) for cell in cells]
            if not any(values[import_indexes[column]] if import_indexes[column] < len(values) else ""
                       for column in IMPORT_COLUMNS):
                continue
            for index in safe_scan_indexes:
                if index >= len(cells):
                    continue
                cell = cells[index]
                if cell.data_type == "f":
                    raise ValueError(f"{row_number}행 {headers[index]}에 수식이 포함되어 있습니다.")
            item = parse_values(values, row_number, headers, safe_scan_indexes, import_indexes, merchant_numbers)
            if item is None:
                continue
            parsed.append(item)
            if len(parsed) > 100_000:
                raise ValueError("한 번에 최대 100,000개 가맹점까지 업로드할 수 있습니다.")

        if not parsed:
            raise ValueError("등록할 가맹점 데이터가 없습니다.")
        return parsed
    finally:
        workbook.close()


def parse_csv(content: bytes) -> list[dict[str, str]]:
    text = None
    for encoding in ("utf-8-sig", "cp949"):
        try:
            text = content.decode(encoding, errors="strict")
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("CSV 인코딩은 UTF-8 또는 CP949여야 합니다.")
    if "\x00" in text:
        raise ValueError("올바른 CSV 파일이 아닙니다.")

    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        headers = [normalize_header(value) for value in next(reader)]
    except StopIteration as exc:
        raise ValueError("CSV 파일이 비어 있습니다.") from exc
    safe_scan_indexes, import_indexes = validate_headers(headers)
    parsed: list[dict[str, str]] = []
    merchant_numbers: set[str] = set()
    try:
        for row_number, raw_values in enumerate(reader, start=2):
            values = [normalize_text(value) for value in raw_values]
            item = parse_values(values, row_number, headers, safe_scan_indexes, import_indexes, merchant_numbers)
            if item is None:
                continue
            parsed.append(item)
            if len(parsed) > 100_000:
                raise ValueError("한 번에 최대 100,000개 가맹점까지 업로드할 수 있습니다.")
    except csv.Error as exc:
        raise ValueError(f"CSV 형식을 읽을 수 없습니다: {exc}") from exc
    if not parsed:
        raise ValueError("등록할 가맹점 데이터가 없습니다.")
    return parsed


async def require_admin(x_admin_key: Annotated[str | None, Header()] = None) -> None:
    configured_key = os.getenv("ADMIN_KEY")
    if not configured_key:
        raise HTTPException(status_code=503, detail="관리자 업로드 키가 설정되지 않았습니다.")
    import secrets
    if not x_admin_key or not secrets.compare_digest(x_admin_key, configured_key):
        raise HTTPException(status_code=401, detail="관리자 인증에 실패했습니다.")


async def kakao_geocode(client: httpx.AsyncClient, address: str) -> tuple[float, float] | None:
    key = os.getenv("KAKAO_REST_API_KEY")
    if not key:
        return None
    try:
        response = await client.get(
            "https://dapi.kakao.com/v2/local/search/address.json",
            headers={"Authorization": f"KakaoAK {key}"},
            params={"query": address, "size": 1},
            timeout=8,
        )
        response.raise_for_status()
        documents = response.json().get("documents", [])
        if not documents:
            return None
        return float(documents[0]["y"]), float(documents[0]["x"])
    except (httpx.HTTPError, KeyError, TypeError, ValueError):
        return None


async def attach_coordinates(rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    with db_connect() as connection:
        existing = {
            row["address"]: (row["lat"], row["lng"])
            for row in connection.execute("SELECT address, lat, lng FROM merchants WHERE lat IS NOT NULL")
        }
    semaphore = asyncio.Semaphore(5)
    async with httpx.AsyncClient() as client:
        async def coordinate(row: dict[str, str]) -> dict[str, Any]:
            result: tuple[float, float] | None = existing.get(row["사업장주소"])
            if result is None:
                async with semaphore:
                    result = await kakao_geocode(client, row["사업장주소"])
            return {**row, "lat": result[0] if result else None, "lng": result[1] if result else None}

        return await asyncio.gather(*(coordinate(row) for row in rows))


def replace_merchants(rows: list[dict[str, Any]], filename: str) -> None:
    now = datetime.now(UTC).isoformat()
    with db_connect() as connection:
        previous_status = {
            row["business_no"]: (row["business_status"], row["status_checked_at"])
            for row in connection.execute("SELECT business_no, business_status, status_checked_at FROM merchants")
        }
        connection.execute("BEGIN IMMEDIATE")
        connection.execute("DELETE FROM merchants")
        connection.executemany(
            """INSERT INTO merchants
               (business_no, merchant_no, name, address, category, lat, lng,
                business_status, status_checked_at, uploaded_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            [
                (
                    row["사업자번호"], row["가맹점번호"], row["가맹점명"],
                    row["사업장주소"], row["가맹점업종명"], row["lat"], row["lng"],
                     *((row["사업자 상태"], now) if row.get("사업자 상태") else
                       previous_status.get(row["사업자번호"], (None, None))), now,
                )
                for row in rows
            ],
        )
        failed = sum(row["lat"] is None for row in rows)
        connection.execute(
            "INSERT INTO upload_logs(filename, row_count, geocode_failed_count, uploaded_at) VALUES (?, ?, ?, ?)",
            (filename, len(rows), failed, now),
        )


def distance_meters(lat1: float, lng1: float, lat2: float, lng2: float) -> float:
    lat_delta = math.radians(lat2 - lat1)
    lng_delta = math.radians(lng2 - lng1)
    value = (
        math.sin(lat_delta / 2) ** 2
        + math.cos(math.radians(lat1)) * math.cos(math.radians(lat2)) * math.sin(lng_delta / 2) ** 2
    )
    return 6_371_000 * 2 * math.atan2(math.sqrt(value), math.sqrt(1 - value))


async def update_business_statuses() -> dict[str, int]:
    service_key = os.getenv("NTS_SERVICE_KEY")
    if not service_key:
        raise HTTPException(status_code=503, detail="국세청 API 키가 설정되지 않았습니다.")
    with db_connect() as connection:
        numbers = [row[0] for row in connection.execute("SELECT DISTINCT business_no FROM merchants")]
    updated = 0
    failed = 0
    async with httpx.AsyncClient(timeout=20) as client:
        for offset in range(0, len(numbers), 100):
            batch = numbers[offset:offset + 100]
            try:
                response = await client.post(
                    "https://api.odcloud.kr/api/nts-businessman/v1/status",
                    params={"serviceKey": service_key, "returnType": "JSON"},
                    json={"b_no": batch},
                )
                response.raise_for_status()
                results = response.json().get("data", [])
                now = datetime.now(UTC).isoformat()
                with db_connect() as connection:
                    connection.executemany(
                        "UPDATE merchants SET business_status = ?, status_checked_at = ? WHERE business_no = ?",
                        [(item.get("b_stt") or "미등록", now, item["b_no"]) for item in results],
                    )
                updated += len(results)
            except (httpx.HTTPError, KeyError, ValueError):
                failed += len(batch)
    return {"updated": updated, "failed": failed}


async def weekly_status_worker() -> None:
    while True:
        await asyncio.sleep(60)
        if not os.getenv("NTS_SERVICE_KEY"):
            await asyncio.sleep(3600)
            continue
        with db_connect() as connection:
            last_checked = connection.execute("SELECT MAX(status_checked_at) FROM merchants").fetchone()[0]
        due = not last_checked or datetime.fromisoformat(last_checked) < datetime.now(UTC) - timedelta(days=7)
        if due:
            try:
                await update_business_statuses()
            except Exception:
                # A transient external API failure must not terminate future weekly checks.
                pass
        await asyncio.sleep(6 * 3600)


@asynccontextmanager
async def lifespan(_: FastAPI):
    initialize_database()
    task = asyncio.create_task(weekly_status_worker())
    yield
    task.cancel()
    try:
        await task
    except asyncio.CancelledError:
        pass


app = FastAPI(
    title="동구랑페이 가맹점 API",
    description="가맹점 지도 검색, 보안 엑셀·CSV 업로드, 사업자 상태 확인 API",
    version="0.1.0",
    lifespan=lifespan,
)
app.add_middleware(SecurityHeadersMiddleware)
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.get("/", include_in_schema=False)
async def index() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/index.html", include_in_schema=False)
async def index_alias() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/styles.css", include_in_schema=False)
async def styles() -> FileResponse:
    return FileResponse(STATIC_DIR / "styles.css", media_type="text/css")


@app.get("/app.js", include_in_schema=False)
async def javascript() -> FileResponse:
    return FileResponse(STATIC_DIR / "app.js", media_type="text/javascript")


@app.get("/merchant-data.js", include_in_schema=False)
async def merchant_data() -> FileResponse:
    return FileResponse(STATIC_DIR / "merchant-data.js", media_type="text/javascript")


@app.get("/donggu-logo.jpg", include_in_schema=False)
async def donggu_logo() -> FileResponse:
    return FileResponse(STATIC_DIR / "donggu-logo.jpg", media_type="image/jpeg")


@app.get("/mobile-preview", include_in_schema=False)
@app.get("/mobile-preview.html", include_in_schema=False)
async def mobile_preview() -> FileResponse:
    return FileResponse(STATIC_DIR / "mobile-preview.html")


@app.get("/admin", include_in_schema=False)
async def admin_page() -> FileResponse:
    return FileResponse(STATIC_DIR / "admin.html")


@app.get("/admin.html", include_in_schema=False)
async def admin_page_alias() -> FileResponse:
    return FileResponse(STATIC_DIR / "admin.html")


@app.get("/health")
async def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/api/config")
async def public_config() -> dict[str, Any]:
    vworld_key = os.getenv("VWORLD_API_KEY", "")
    with db_connect() as connection:
        has_uploaded_map = bool(connection.execute(
            "SELECT EXISTS(SELECT 1 FROM upload_logs) AND EXISTS(SELECT 1 FROM merchants WHERE lat IS NOT NULL)"
        ).fetchone()[0])
    return {
        "tile_provider": "vworld" if vworld_key else "vworld_public",
        "vworld_key": vworld_key,
        "kakao_geocoding": bool(os.getenv("KAKAO_REST_API_KEY")),
        "demo": not has_uploaded_map,
    }


@app.get("/api/merchants", response_model=MerchantList)
async def list_merchants(
    q: str | None = Query(None, max_length=100),
    category: str | None = Query(None, max_length=100),
    lat: float | None = Query(None, ge=-90, le=90),
    lng: float | None = Query(None, ge=-180, le=180),
    radius: int = Query(1000, ge=100, le=20_000),
    south: float | None = Query(None, ge=-90, le=90),
    west: float | None = Query(None, ge=-180, le=180),
    north: float | None = Query(None, ge=-90, le=90),
    east: float | None = Query(None, ge=-180, le=180),
    limit: int = Query(5000, ge=1, le=10_000),
) -> MerchantList:
    if (lat is None) != (lng is None):
        raise HTTPException(status_code=422, detail="lat와 lng는 함께 입력해야 합니다.")
    bounds = (south, west, north, east)
    if any(value is not None for value in bounds) and any(value is None for value in bounds):
        raise HTTPException(status_code=422, detail="south, west, north, east는 함께 입력해야 합니다.")
    clauses = ["lat IS NOT NULL", "lng IS NOT NULL", "COALESCE(business_status, '') NOT LIKE '%폐업%'"]
    parameters: list[Any] = []
    if q:
        clauses.append("(name LIKE ? ESCAPE '\\' OR address LIKE ? ESCAPE '\\' OR category LIKE ? ESCAPE '\\')")
        escaped = q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        parameters.extend([f"%{escaped}%", f"%{escaped}%", f"%{escaped}%"])
    if category:
        clauses.append("category = ?")
        parameters.append(category)
    if None not in (south, west, north, east):
        clauses.extend(["lat BETWEEN ? AND ?", "lng BETWEEN ? AND ?"])
        parameters.extend([south, north, west, east])
    if lat is not None and lng is not None:
        lat_delta = radius / 111_320
        lng_delta = radius / max(1, 111_320 * math.cos(math.radians(lat)))
        clauses.extend(["lat BETWEEN ? AND ?", "lng BETWEEN ? AND ?"])
        parameters.extend([lat - lat_delta, lat + lat_delta, lng - lng_delta, lng + lng_delta])
    sql = f"SELECT * FROM merchants WHERE {' AND '.join(clauses)} LIMIT ?"
    parameters.append(limit)
    with db_connect() as connection:
        rows = connection.execute(sql, parameters).fetchall()

    items: list[Merchant] = []
    for row in rows:
        distance = distance_meters(lat, lng, row["lat"], row["lng"]) if lat is not None and lng is not None else None
        if distance is not None and distance > radius:
            continue
        items.append(Merchant(
            id=row["id"], name=row["name"], address=row["address"], category=row["category"],
            lat=row["lat"], lng=row["lng"], distance=distance, business_status=row["business_status"],
        ))
    if lat is not None and lng is not None:
        items.sort(key=lambda item: item.distance if item.distance is not None else math.inf)
    return MerchantList(items=items, total=len(items))


@app.get("/api/categories")
async def categories() -> dict[str, list[str]]:
    with db_connect() as connection:
        values = [row[0] for row in connection.execute(
            "SELECT DISTINCT category FROM merchants WHERE COALESCE(business_status, '') NOT LIKE '%폐업%' ORDER BY category"
        )]
    return {"items": values}


@app.get("/api/geocode")
async def geocode(query: str = Query(..., min_length=2, max_length=200)) -> dict[str, list[dict[str, Any]]]:
    async with httpx.AsyncClient() as client:
        result = await kakao_geocode(client, query)
    if not result:
        return {"items": []}
    return {"items": [{"lat": result[0], "lng": result[1], "address": query}]}


@app.post("/api/admin/upload", dependencies=[Depends(require_admin)])
async def upload_merchants(file: Annotated[UploadFile, File(...)]) -> dict[str, Any]:
    filename = Path(file.filename or "").name
    extension = Path(filename).suffix.lower()
    if extension not in {".xlsx", ".csv"}:
        raise HTTPException(status_code=400, detail=".xlsx 또는 .csv 파일만 업로드할 수 있습니다.")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="파일 크기는 10MB를 초과할 수 없습니다.")
    try:
        parser = parse_excel if extension == ".xlsx" else parse_csv
        rows = await run_in_threadpool(parser, content)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    coordinated_rows = await attach_coordinates(rows)
    await run_in_threadpool(replace_merchants, coordinated_rows, filename)
    failed = sum(row["lat"] is None for row in coordinated_rows)
    return {
        "message": "가맹점 데이터가 교체되었습니다.",
        "imported": len(coordinated_rows),
        "geocoded": len(coordinated_rows) - failed,
        "geocode_failed": failed,
        "discarded_columns": sorted(DISCARDED_PII_COLUMNS),
    }


@app.post("/api/admin/check-business-status", dependencies=[Depends(require_admin)])
async def check_business_status() -> dict[str, int]:
    return await update_business_statuses()
